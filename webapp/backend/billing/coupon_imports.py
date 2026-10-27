"""Review original KMS A/K spreadsheet observations before applying them."""
from datetime import date
from decimal import Decimal, InvalidOperation
import hashlib
from io import BytesIO
from pathlib import PurePath
import re
from uuid import uuid4
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from models import Student
from .coupon_accounting import issue, lock_student, operation, require_writes, save_operation, storage_enabled, writes_enabled
from .coupon_models import CouponAccount, CouponImport, CouponObservation
from .coupon_schemas import CouponImportReview, CouponImportRow
from .drafts import DraftConflict, DraftNotFound, digest
from .models import utc_now
from .schemas import BillingStudent

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000
COMPANY_ID = re.compile(r'(MSA|MSB)([0-9]+)\Z', re.ASCII)
KMS_HEADERS = ['ID#', 'Student', 'Grade', 'Instructor', 'Schedule', 'LastLesson', 'PaidMth',
               '60DAbs', 'Remain', 'Balance', 'Coupon', 'RegularFee', 'Parent', 'Mobile']


def _count(value):
    if value is None or (isinstance(value, str) and value.strip() in {'', '--'}):
        return 0
    if isinstance(value, bool):
        raise ValueError('Check the coupon count in column K.')
    try:
        number = Decimal(str(value).strip())
        if not number.is_finite() or number != number.to_integral_value() or not 0 <= number <= 10000:
            raise ValueError('Check the coupon count in column K.')
        return int(number)
    except (InvalidOperation, ValueError):
        raise ValueError('Check the coupon count in column K.') from None


def parse_workbook(content):
    """Read text IDs exactly. Invalid counts remain invalid, never zero."""
    if not isinstance(content, bytes) or not content or len(content) > MAX_BYTES:
        raise ValueError('Choose a KMS spreadsheet smaller than 5 MB.')
    if content.startswith(bytes.fromhex('d0cf11e0a1b11ae1')):
        import xlrd
        try:
            workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
            sheet = workbook.sheet_by_index(0)
            if sheet.ncols < 11 or sheet.nrows > MAX_ROWS + 1:
                raise ValueError('Choose a KMS list with columns A to K and at most 10,000 records.')
            raw_rows = [(sheet.row_values(i), False)
                        for i in range(sheet.nrows)]
            return _parsed_rows(raw_rows)
        except xlrd.XLRDError:
            raise ValueError('Choose the original KMS Excel spreadsheet.') from None
        finally:
            if 'workbook' in locals():
                workbook.release_resources()
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError('The KMS spreadsheet is too large to inspect safely.')
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except BadZipFile:
        raise ValueError('Save this KMS list as an Excel .xlsx file before importing it.') from None
    except (KeyError, OSError, InvalidFileException):
        raise ValueError('Choose the original KMS spreadsheet.') from None
    try:
        if not workbook.worksheets:
            raise ValueError('The spreadsheet has no worksheet.')
        sheet = workbook.worksheets[0]
        if sheet.max_column < 11 or sheet.max_row > MAX_ROWS + 1:
            raise ValueError('Choose a KMS list with columns A to K and at most 10,000 records.')
        return _parsed_rows((([cell.value for cell in cells],
                             cells[0].data_type == 'f' or cells[10].data_type == 'f'))
                            for cells in sheet.iter_rows(min_col=1, max_col=min(sheet.max_column, 14)))
    finally:
        workbook.close()


def _parsed_rows(raw_rows):
    raw_rows = list(raw_rows)
    metadata = set()
    if len(raw_rows) >= 3:
        title, _ = raw_rows[0]
        headers, _ = raw_rows[1]
        if (re.fullmatch(r'MS[AB] / \d{4}-\d{2}-\d{2}', str(title[0] or '').strip())
                and [str(value or '').strip() for value in headers[:14]] == KMS_HEADERS):
            metadata.update({0, 1})
            footer, formula = raw_rows[-1]
            if (len(footer) >= 14 and not formula and all(not str(value or '').strip() for value in footer[:7])
                    and all(not str(value or '').strip() for value in footer[12:14])
                    and str(footer[9]).strip() == '--' and str(footer[11]).strip() == '--'
                    and all(re.fullmatch(r'[0-9]+(?:,[0-9]{3})*', str(footer[i]).strip()) for i in (7, 8))):
                try:
                    total = sum(_count(values[10]) for values, _ in raw_rows[2:-1]
                                if COMPANY_ID.fullmatch(str(values[0] or '').strip()))
                    if _count(footer[10]) == total:
                        metadata.add(len(raw_rows) - 1)
                except ValueError:
                    pass  # Invalid student counts remain visible for review.
    result = []
    for number, (values, formula) in enumerate(raw_rows, start=1):
        company, raw_count = values[0], values[10]
        if number - 1 in metadata:
            continue
        if number == 1 and str(company or '').strip().casefold() in {'id#', 'id', 'student id', 'company id'}:
            continue
        if str(company or '').strip() == '' and (raw_count is None or str(raw_count).strip() == ''):
            continue
        record = {'row_number': number, 'company_id': str(company or '').strip(), 'remaining_count': None, 'error': None}
        try:
            if formula or not COMPANY_ID.fullmatch(record['company_id']):
                raise ValueError('Check the student ID in column A.')
            record['remaining_count'] = _count(raw_count)
        except ValueError as error:
            record['error'] = str(error)
        result.append(record)
    if not result:
        raise ValueError('The KMS spreadsheet contains no student records.')
    return result


def _students(db, branch, company_ids):
    school_ids = {match.group(2) for value in company_ids if (match := COMPANY_ID.fullmatch(value)) and match.group(1) == branch}
    records = {}
    for row in db.query(Student.id, Student.student_name, Student.school_student_id, Student.home_location,
                        Student.grade, Student.is_staff_referral, Student.lang_stream, Student.school).filter(
            Student.home_location == branch, Student.school_student_id.in_(school_ids)).all():
        records.setdefault(branch + row.school_student_id, []).append(BillingStudent(**{
            **dict(row._mapping), 'is_staff_referral': bool(row.is_staff_referral)}))
    return records


def review_import(db, imported, *, as_of):
    if not storage_enabled():
        raise DraftConflict('Coupon imports are not available yet.')
    if imported.applied_review:
        return CouponImportReview.model_validate(imported.applied_review)
    records = _students(db, imported.branch, [r['company_id'] for r in imported.parsed_rows])
    ids = {s.id for values in records.values() for s in values}
    latest = {sid: observed for sid, observed in db.query(CouponAccount.student_id, CouponObservation).join(
        CouponObservation, CouponObservation.id == CouponAccount.observation_id).filter(CouponAccount.student_id.in_(ids)).all()} if ids else {}
    versions = dict(db.query(CouponAccount.student_id, CouponAccount.version).filter(CouponAccount.student_id.in_(ids)).all()) if ids else {}
    errors = []
    if imported.observed_on > as_of:
        errors.append(issue('FUTURE_OBSERVATION', 'The KMS observation date cannot be in the future.'))
    conflicting_sha = db.query(CouponImport.id).filter(CouponImport.source_sha256 == imported.source_sha256,
        CouponImport.id != imported.id, CouponImport.observed_on != imported.observed_on,
        CouponImport.applied_at.is_not(None)).first()
    if conflicting_sha:
        errors.append(issue('SOURCE_DATE_CHANGED', 'This file was already imported with another observation date. Check the original KMS date.'))
    seen = set()
    rows = []
    for raw in imported.parsed_rows:
        company = raw['company_id']
        row = CouponImportRow(row_number=raw['row_number'], company_id=company,
                              remaining_count=raw['remaining_count'], status='invalid')
        match = COMPANY_ID.fullmatch(company)
        if raw['error'] or not match or match.group(1) != imported.branch:
            row.issues.append(issue('INVALID_KMS_ROW', raw['error'] or 'Choose the matching KMS profile branch.'))
        elif company in seen:
            row.issues.append(issue('DUPLICATE_STUDENT', 'The spreadsheet repeats this student ID.'))
        elif len(records.get(company, [])) > 1:
            row.issues.append(issue('AMBIGUOUS_STUDENT', 'More than one student has this profile branch and ID.'))
        elif not records.get(company):
            row.status = 'unmatched'
            row.issues.append(issue('STUDENT_NOT_FOUND', 'No matching student profile was found.'))
        else:
            row.student = records[company][0]
            previous = latest.get(row.student.id)
            row.previous_count = previous.remaining_count if previous else None
            if previous and previous.observed_on > imported.observed_on:
                row.issues.append(issue('OLDER_OBSERVATION', 'A newer KMS observation is already recorded.'))
            else:
                row.status = 'unchanged' if row.previous_count == row.remaining_count else 'changed'
                if previous and previous.observed_on == imported.observed_on and previous.source_sha256 != imported.source_sha256:
                    row.issues.append(issue('SAME_DAY_SOURCE_REVIEW', 'Confirm that this is the latest KMS export for this date.'))
        seen.add(company)
        rows.append(row)
    counts = {status: sum(r.status == status for r in rows) for status in ('changed', 'unchanged', 'unmatched', 'invalid')}
    payload = {'id': imported.id, 'branch': imported.branch, 'observed_on': imported.observed_on,
        'source_sha256': imported.source_sha256, 'rows': [r.model_dump(mode='json') for r in rows],
        'versions': versions, 'observations': {sid: r.id for sid, r in latest.items()},
        'issues': [i.model_dump() for i in errors]}
    return CouponImportReview(id=imported.id, branch=imported.branch, observed_on=imported.observed_on,
        filename=imported.source_filename, source_sha256=imported.source_sha256, review_hash=digest(payload),
        changed_count=counts['changed'], unchanged_count=counts['unchanged'], unmatched_count=counts['unmatched'],
        invalid_count=counts['invalid'], rows=rows, issues=errors, applied=False,
        can_confirm=writes_enabled() and not errors and not counts['invalid'] and bool(counts['changed'] + counts['unchanged']))


def preview_import(db, content, filename, branch, observed_on, actor, *, as_of):
    require_writes()
    if branch not in {'MSA', 'MSB'}:
        raise ValueError('Choose profile branch MSA or MSB.')
    rows = parse_workbook(content)
    sha = hashlib.sha256(content).hexdigest()
    existing = db.query(CouponImport).filter_by(branch=branch, observed_on=observed_on, source_sha256=sha).first()
    if existing:
        return review_import(db, existing, as_of=as_of)
    imported = CouponImport(id=str(uuid4()), branch=branch, observed_on=observed_on,
        source_sha256=sha, source_filename=PurePath(filename.replace('\\', '/')).name[:500],
        parsed_rows=rows, created_by_name=(actor.tutor_name or str(actor.id))[:100])
    db.add(imported)
    db.flush()
    return review_import(db, imported, as_of=as_of)


def find_import(db, import_id):
    if not storage_enabled():
        raise DraftConflict('Coupon imports are not available yet.')
    row = db.query(CouponImport).filter_by(id=str(import_id)).first()
    if row is None:
        raise DraftNotFound('Coupon import not found.')
    return row


def apply_import(db, import_id, data, actor, *, as_of):
    require_writes()
    initial = find_import(db, import_id)
    records = _students(db, initial.branch, [r['company_id'] for r in initial.parsed_rows])
    ids = sorted({s.id for values in records.values() for s in values})
    # Restart after ownership discovery so MySQL's consistent read begins only
    # after all affected Student locks, in the same order as financial actions.
    db.rollback()
    if ids:
        db.query(Student.id).filter(Student.id.in_(ids)).order_by(Student.id).with_for_update().all()
    imported = db.query(CouponImport).filter_by(id=str(import_id)).populate_existing().with_for_update().one()
    payload_hash = digest({'kind': 'import_applied', 'import_id': str(import_id), 'data': data.model_dump(mode='json')})
    replay = operation(db, actor, data.idempotency_key, payload_hash, CouponImportReview)
    if replay:
        return replay
    if imported.applied_review:
        return CouponImportReview.model_validate(imported.applied_review)
    current = review_import(db, imported, as_of=as_of)
    if current.review_hash != data.expected_review_hash:
        raise DraftConflict('The KMS import review changed. Refresh it before applying.')
    if not current.can_confirm or (current.unmatched_count and not data.acknowledge_unmatched):
        raise DraftConflict('Resolve invalid rows and acknowledge unmatched students before applying this list.')
    if any(i.code == 'SAME_DAY_SOURCE_REVIEW' for r in current.rows for i in r.issues) and not data.acknowledge_same_day:
        raise DraftConflict('Confirm that this is the latest KMS export for this date before replacing the current observation.')
    for row in current.rows:
        if row.student is None:
            continue
        account = lock_student(db, row.student.id)
        observation = CouponObservation(id=str(uuid4()), student_id=row.student.id, import_id=imported.id,
            remaining_count=row.remaining_count, observed_on=imported.observed_on,
            source_filename=imported.source_filename, source_sha256=imported.source_sha256)
        db.add(observation)
        db.flush()
        account.observation_id, account.verification_id = observation.id, None
        account.version += 1
    # Completion describes the checks still needed after this exact import.
    # Read all affected students together, without one request per student.
    from .coupon_accounting import balances_for_students
    follow_up = balances_for_students(db, ids, as_of=as_of)
    for row in current.rows:
        if row.student is not None:
            row.coupon_check_required = follow_up[row.student.id].usable_count is None
    imported.applied_at = utc_now()
    current.applied, current.can_confirm = True, False
    imported.applied_review = current.model_dump(mode='json')
    db.flush()
    save_operation(db, actor, data.idempotency_key, payload_hash, 'import_applied', current, import_id=imported.id)
    return current

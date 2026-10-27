"""Generate a reviewed bank workbook without submitting or settling payments.

The bank identifies contracts through the student's profile branch and exact
stored school number. Teaching-branch filters do not define a bank account.
"""
from collections import Counter
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import Workbook

from .autopay import build_autopay_run
from .bank_export_schemas import BankExportIssue, BankExportPreview, BankExportRow
from .calculator import BillingPolicy
from .drafts import digest, utc
from .models import BillingAccount, BillingPayment, BillingPreference
from .workflow import workflow_storage_enabled

HEADERS = ('Contract No', 'Amount', 'Remark')
SHEET_TITLE = 'Batch_DirectDebit_Sample'
FORMAT_VERSION = 'contract-amount-remark-v1'
CONTRACT_PATTERN = re.compile(r'[A-Za-z1-9][A-Za-z0-9-]{0,19}', flags=re.ASCII)
MONEY_UNIT = Decimal('0.01')
BANK_AMOUNT_LIMIT = Decimal('100000000000')


class ExportConflict(ValueError):
    pass


class ExportBlocked(ValueError):
    pass


class ExportUnavailable(ValueError):
    pass


def issue(code, message, student_id=None):
    return BankExportIssue(code=code, message=message, student_id=student_id)


def contract_number(student):
    if student.home_location not in {'MSA', 'MSB'}:
        raise ValueError('Set the student’s profile branch to MSA or MSB.')
    school_id = student.school_student_id
    if not isinstance(school_id, str) or not school_id:
        raise ValueError('Enter the student’s school student ID.')
    contract = student.home_location + '0' + school_id
    if not CONTRACT_PATTERN.fullmatch(contract):
        raise ValueError('Check the school student ID. The contract must use English letters, numbers or hyphens and fit within 20 characters.')
    return contract


def bank_amount(value):
    """Validate the exact value without rounding an unsupported amount."""
    if value is None or isinstance(value, bool):
        raise ValueError('Check the confirmed payment amount.')
    try:
        amount = Decimal(str(value))
        if (not amount.is_finite() or not Decimal('0') < amount < BANK_AMOUNT_LIMIT
                or amount != amount.quantize(MONEY_UNIT)):
            raise ValueError('Check the confirmed payment amount.')
        return amount.quantize(MONEY_UNIT)
    except (InvalidOperation, TypeError, ValueError):
        raise ValueError('Check the confirmed payment amount.') from None


def _payment_state(db, ids):
    """Include preference revisions and frozen teaching terms in the review hash."""
    if not ids:
        return []
    result = []
    for row in db.query(BillingPayment.id, BillingPayment.version, BillingPayment.status,
            BillingPayment.snapshot, BillingPayment.sent_at, BillingAccount.student_id,
            BillingAccount.active, BillingPreference.version.label('preference_version'),
            BillingPreference.payment_method, BillingPreference.mandate_ready).join(
            BillingAccount, BillingAccount.id == BillingPayment.account_id).outerjoin(
            BillingPreference, BillingPreference.account_id == BillingAccount.id).filter(
            BillingPayment.id.in_(ids)).order_by(BillingPayment.id).all():
        result.append({'payment_id': row.id, 'version': row.version, 'status': row.status,
            'snapshot_hash': digest(row.snapshot), 'sent_at': utc(row.sent_at).isoformat() if row.sent_at else None,
            'student_id': row.student_id, 'account_active': bool(row.active),
            'preference_version': row.preference_version, 'payment_method': row.payment_method,
            'mandate_ready': bool(row.mandate_ready)})
    return result


def preview_export(db, *, as_of: date, month: date, branch: str, holidays: set[date],
                   policy=BillingPolicy()):
    """Use every current Ready payment for the profile branch, without UI filters."""
    from routers.billing import student_number_order

    issues = []
    rows = []
    ready = []
    payment_state = []
    claim_context = {}
    excluded_claimed_count = 0
    all_amounts_valid = True
    if not workflow_storage_enabled():
        issues.append(issue('CONFIRMED_STORAGE_UNAVAILABLE', 'Confirmed Autopay payments are not available yet.'))
    else:
        # This is deliberately unfiltered by teaching location, search, tab,
        # page or viewed tutor. The endpoints require effective Admin access.
        current = build_autopay_run(db, as_of=as_of, month=month, holidays=holidays, policy=policy)
        from .bank_runs import claimed_payments, storage_enabled as run_storage_enabled
        claims = claimed_payments(db, [r.payment_id for r in current if r.stage == 'ready'])
        for row in current:
            if row.stage != 'ready':
                continue
            if row.payment_id in claims:
                if row.student.home_location == branch:
                    excluded_claimed_count += 1
                    claim_context[row.payment_id] = claims[row.payment_id]
                continue
            if row.student.home_location not in {'MSA', 'MSB'}:
                issues.append(issue('PROFILE_BRANCH_REQUIRED', 'Set the student’s profile branch to MSA or MSB.', row.student.id))
            elif row.student.home_location == branch:
                ready.append(row)
        ready.sort(key=lambda r: (student_number_order(r.student), r.payment_id or ''))
        for row in ready:
            contract = ''
            amount = None
            try:
                contract = contract_number(row.student)
            except ValueError as error:
                issues.append(issue('INVALID_CONTRACT', str(error), row.student.id))
            try:
                amount = format(bank_amount(row.total_fee), '.2f')
            except ValueError as error:
                all_amounts_valid = False
                issues.append(issue('INVALID_AMOUNT', str(error), row.student.id))
            # Defend this boundary as well as the run classification. A draft
            # or projection can never become a downloadable bank row.
            if (row.origin != 'payment' or not row.payment_id or not row.payment_version
                    or row.payment_status != 'Pending Payment' or row.amount_basis != 'confirmed'):
                issues.append(issue('CONFIRMED_PAYMENT_REQUIRED', 'Refresh the confirmed Autopay payment.', row.student.id))
                continue
            rows.append(BankExportRow(payment_id=row.payment_id, payment_version=row.payment_version,
                                      student=row.student, contract_no=contract, amount=amount))
        duplicates = {contract for contract, count in Counter(r.contract_no for r in rows if r.contract_no).items() if count > 1}
        for row in rows:
            if row.contract_no in duplicates:
                issues.append(issue('DUPLICATE_CONTRACT', 'Two payments use the same contract number. Check the student IDs.', row.student.id))
        if not ready:
            issues.append(issue('NO_READY_PAYMENTS', 'No Ready Autopay payments for this profile branch.'))
        payment_state = _payment_state(db, [r.payment_id for r in rows])
        if {p['payment_id'] for p in payment_state} != {r.payment_id for r in rows}:
            issues.append(issue('PAYMENT_CHANGED', 'The Ready payments changed. Refresh the export review.'))
    # This is the only total calculation. The validated output rows feed both
    # the numeric workbook cells and the response header on download.
    total = sum((Decimal(row.amount) for row in rows), Decimal('0.00')) if all_amounts_valid else None
    total_amount = format(total, '.2f') if total is not None else None
    issues.sort(key=lambda i: (i.code, i.student_id or 0, i.message))
    payload = {'format': FORMAT_VERSION, 'collection_month': month.strftime('%Y-%m'), 'branch': branch,
        'rows': [r.model_dump(mode='json') for r in rows], 'total_amount': total_amount,
        'readiness': [r.model_dump(mode='json', exclude={'bank_run'}) for r in ready], 'payment_state': payment_state,
        'issues': [i.model_dump(mode='json') for i in issues]}
    if workflow_storage_enabled():
        from .bank_runs import storage_enabled as run_storage_enabled
        if run_storage_enabled():
            payload['claimed_payments'] = claim_context
    content_hash = digest(payload)
    return BankExportPreview(collection_month=month.strftime('%Y-%m'), branch=branch,
        payment_count=len(rows), excluded_claimed_count=excluded_claimed_count, total_amount=total_amount, content_hash=content_hash,
        filename=f'autopay-{branch}-{month:%Y-%m}-{content_hash[:12]}.xlsx',
        rows=rows, issues=issues, can_download=bool(rows) and not issues)


def workbook_bytes(preview: BankExportPreview):
    if not preview.can_download:
        raise ExportBlocked('Resolve the export checks before downloading.')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE
    sheet.append(HEADERS)
    sheet.column_dimensions['A'].number_format = '@'
    sheet.column_dimensions['B'].number_format = 'General'
    sheet.column_dimensions['C'].number_format = '@'
    for index, row in enumerate(preview.rows, start=2):
        contract = sheet.cell(index, 1, row.contract_no)
        contract.data_type, contract.number_format = 's', '@'
        amount = sheet.cell(index, 2, Decimal(row.amount))
        amount.number_format = 'General'
        remark = sheet.cell(index, 3)
        remark.number_format = '@'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    # openpyxl formats numeric values through binary floats. At the bank's
    # upper limit, 99999999999.99 would become 99999999999.99001 in its XML.
    # Keep numeric cells, but write the validated decimal literals exactly.
    sheet_path = 'xl/worksheets/sheet1.xml'
    namespace = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    corrected = BytesIO()
    with ZipFile(BytesIO(output.getvalue())) as source, ZipFile(corrected, 'w') as target:
        for member in source.infolist():
            content = source.read(member.filename)
            if member.filename == sheet_path:
                root = ElementTree.fromstring(content)
                cells = {cell.attrib['r']: cell for cell in root.findall('m:sheetData/m:row/m:c', namespace)}
                for index, row in enumerate(preview.rows, start=2):
                    cell = cells[f'B{index}']
                    cell.set('t', 'n')
                    cell.find('m:v', namespace).text = row.amount
                content = ElementTree.tostring(root, encoding='utf-8', xml_declaration=True)
            target.writestr(member, content)
    return corrected.getvalue()


def reviewed_export(db, *, expected_content_hash, **kwargs):
    if not workflow_storage_enabled():
        raise ExportUnavailable('Confirmed Autopay payments are not available yet.')
    preview = preview_export(db, **kwargs)
    if preview.content_hash != expected_content_hash:
        raise ExportConflict('The export changed. Refresh and review it before downloading.')
    if not preview.can_download:
        raise ExportBlocked('Resolve the export checks before downloading.')
    return preview, workbook_bytes(preview)

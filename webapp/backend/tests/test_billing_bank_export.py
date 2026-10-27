"""Read-only bank files from reviewed Ready records in disposable SQLite."""
from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from sqlalchemy import event
from sqlalchemy.engine import Engine

from billing.bank_export import bank_amount, contract_number
from billing.models import BillingOperation, BillingPayment, BillingPreference
from billing.schemas import BillingStudent
from models import Student
from tests.test_billing_drafts import env as draft_env
from tests.test_billing_workflow import env as workflow_env
from tests.test_billing_autopay import env as autopay_env, autopay_payment, clone_payment_for_student


@pytest.fixture(autouse=True)
def reject_external_connections():
    def reject(dialect, *_args):
        assert dialect.name == 'sqlite', 'Bank export checks only permit disposable SQLite.'
    event.listen(Engine, 'do_connect', reject)
    yield
    event.remove(Engine, 'do_connect', reject)


@pytest.fixture
def env(autopay_env):
    client, db, user = autopay_env
    student = db.get(Student, 101)
    student.home_location, student.school_student_id = 'MSA', '2100'
    db.commit()
    payment = autopay_payment(client, db)
    return client, db, user, payment


def preview(client, branch='MSA', **changes):
    body = {'collection_month': '2026-11', 'branch': branch, **changes}
    response = client.post('/api/billing/autopay/export-preview', json=body)
    assert response.status_code == 200, response.text
    return response.json()


def download(client, review, **changes):
    return client.post('/api/billing/autopay/export', json={
        'collection_month': review['collection_month'], 'branch': review['branch'],
        'expected_content_hash': review['content_hash'], **changes})


def test_exact_workbook_headers_numeric_amounts_and_download_total_are_query_only(env):
    client, db, _, payment = env
    operations = db.query(BillingOperation).count()
    statements = []
    def select_only(_conn, _cursor, sql, *_args):
        assert sql.lstrip().upper().startswith('SELECT'), sql
        statements.append(sql)
    event.listen(db.bind, 'before_cursor_execute', select_only)
    try:
        review = preview(client)
        response = download(client, review)
    finally:
        event.remove(db.bind, 'before_cursor_execute', select_only)
    assert response.status_code == 200, response.text
    assert review['can_download'] and not review['issues']
    assert review['rows'][0]['payment_id'] == payment.id
    assert review['rows'][0]['contract_no'] == 'MSA02100'
    assert review['rows'][0]['amount'] == review['total_amount']
    assert response.headers['x-csm-export-total'] == review['total_amount']
    assert response.headers['x-csm-export-count'] == str(review['payment_count']) == '1'
    assert response.headers['x-csm-export-hash'] == review['content_hash']
    assert response.headers['content-disposition'] == f'attachment; filename="{review["filename"]}"'
    assert review['filename'].endswith('.xlsx') and response.content[:2] == b'PK'
    assert response.headers['cache-control'] == 'no-store'
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ['Batch_DirectDebit_Sample']
    sheet = workbook.active
    assert list(sheet.values) == [('Contract No', 'Amount', 'Remark'),
                                  ('MSA02100', int(review['total_amount'].split('.')[0]), None)]
    assert sheet['A2'].data_type == 's' and sheet['A2'].number_format == '@'
    assert sheet['B2'].data_type == 'n' and sheet['B2'].number_format == 'General'
    assert sheet['C2'].value is None and sheet['C2'].number_format == '@'
    with ZipFile(BytesIO(response.content)) as archive:
        xml = archive.read('xl/worksheets/sheet1.xml').decode()
        assert 'Synthetic' not in xml and '<f>' not in xml and 'MSA02100' in xml
    assert statements and db.query(BillingOperation).count() == operations
    assert db.get(BillingPayment, payment.id).status == 'Pending Payment'
    assert not db.dirty and not db.new


def test_profile_branch_groups_all_ready_independent_of_list_filters(env):
    client, db, _, payment = env
    other = clone_payment_for_student(db, payment, 201)
    db.get(Student, 201).home_location = 'MSB'
    db.get(Student, 201).school_student_id = '0011'
    db.commit()
    msa, msb = preview(client), preview(client, 'MSB')
    assert msa['can_download'] and msb['can_download']
    assert {row['payment_id'] for row in msa['rows']} == {payment.id}
    assert {row['payment_id'] for row in msb['rows']} == {other.id}
    assert msb['rows'][0]['contract_no'] == 'MSB00011'
    # Both classes were taught at MSA. Profile branch alone chooses the file.
    filtered = client.post('/api/billing/autopay/export-preview?location=MSB&search=missing&offset=50&limit=1&stage=skipped',
                          json={'collection_month': '2026-11', 'branch': 'MSA'})
    assert filtered.status_code == 200 and filtered.json()['content_hash'] == msa['content_hash']
    assert len(load_workbook(BytesIO(download(client, msb).content)).active['A']) == 2


@pytest.mark.parametrize('school_id', ['', '=1+1', '+123', '@ID', '1\n2', '中1', ' 123', '123.0', '1' * 17])
def test_invalid_contract_blocks_the_whole_file_without_dropping_the_student(env, school_id):
    client, db, _, _ = env
    db.get(Student, 101).school_student_id = school_id
    db.commit()
    review = preview(client)
    assert not review['can_download']
    assert any(issue['code'] == 'INVALID_CONTRACT' for issue in review['issues'])
    assert review['payment_count'] == len(review['rows']) == 1
    assert review['rows'][0]['contract_no'] == '' and review['total_amount'] != '0.00'
    assert download(client, review).status_code == 422


def test_missing_profile_branch_blocks_both_files_but_other_known_branch_invalid_id_does_not(env):
    client, db, _, payment = env
    clone_payment_for_student(db, payment, 201)
    other = db.get(Student, 201)
    other.home_location, other.school_student_id = 'MSB', ''
    db.commit()
    assert preview(client)['can_download']
    assert not preview(client, 'MSB')['can_download']
    other.home_location = None
    db.commit()
    for branch in ('MSA', 'MSB'):
        review = preview(client, branch)
        assert not review['can_download']
        assert any(i['code'] == 'PROFILE_BRANCH_REQUIRED' and i['student_id'] == 201 for i in review['issues'])


def test_duplicate_contracts_block_without_combining_amounts(env):
    client, db, _, payment = env
    clone_payment_for_student(db, payment, 201)
    db.get(Student, 201).school_student_id = '2100'
    db.commit()
    review = preview(client)
    assert not review['can_download'] and review['payment_count'] == 2
    assert {i['student_id'] for i in review['issues'] if i['code'] == 'DUPLICATE_CONTRACT'} == {101, 201}
    assert Decimal(review['total_amount']) == sum(Decimal(r['amount']) for r in review['rows'])
    assert download(client, review).status_code == 422


@pytest.mark.parametrize('change', ['contract', 'branch', 'preference_revision', 'lost_ready', 'new_ready', 'payment_version'])
def test_download_requires_the_same_reviewed_set_and_state(env, change):
    client, db, _, payment = env
    review = preview(client)
    if change == 'contract': db.get(Student, 101).school_student_id = '2101'
    if change == 'branch': db.get(Student, 101).home_location = 'MSB'
    if change == 'preference_revision': db.query(BillingPreference).one().version += 1
    if change == 'lost_ready': db.query(BillingPreference).one().mandate_ready = False
    if change == 'new_ready': clone_payment_for_student(db, payment, 201)
    if change == 'payment_version': payment.version += 1
    db.commit()
    response = download(client, review)
    assert response.status_code == 409, response.text
    assert 'refresh' in response.json()['detail'].lower()


@pytest.mark.parametrize('role,header', [('Tutor', None), ('Supervisor', None), ('Guest', None),
    ('Super Admin', 'Tutor'), ('Super Admin', 'Supervisor'), ('Super Admin', 'Guest')])
def test_export_role_boundary_includes_impersonation(env, role, header):
    client, _, user, _ = env
    user.role = role
    headers = {'X-Effective-Role': header} if header else {}
    for route, body in [('/api/billing/autopay/export-preview', {'collection_month': '2026-11', 'branch': 'MSA'}),
        ('/api/billing/autopay/export', {'collection_month': '2026-11', 'branch': 'MSA', 'expected_content_hash': 'a' * 64})]:
        assert client.post(route, json=body, headers=headers).status_code == 403


def test_storage_reads_work_independently_of_writes_but_187_does_not_query_188(env, monkeypatch):
    client, db, _, _ = env
    monkeypatch.setenv('BILLING_WORKFLOW_ENABLED', 'false')
    monkeypatch.setenv('BILLING_WORKFLOW_STORAGE_ENABLED', 'true')
    review = preview(client)
    assert review['can_download'] and download(client, review).status_code == 200
    monkeypatch.setenv('BILLING_WORKFLOW_STORAGE_ENABLED', 'false')
    def old_schema_only(_conn, _cursor, sql, *_args):
        assert not any(name in sql.lower() for name in ('billing_preferences', 'billing_payments',
            'billing_payment_lines', 'billing_lesson_claims', 'billing_operations')), sql
    event.listen(db.bind, 'before_cursor_execute', old_schema_only)
    try:
        blocked = preview(client)
        assert not blocked['can_download'] and blocked['payment_count'] == 0
        assert blocked['total_amount'] == '0.00'
        assert blocked['issues'][0]['code'] == 'CONFIRMED_STORAGE_UNAVAILABLE'
        assert download(client, blocked).status_code == 503
    finally:
        event.remove(db.bind, 'before_cursor_execute', old_schema_only)


@pytest.mark.parametrize('value', ['NaN', 'Infinity', '-1', '0', '1.001', '100000000000', None, True])
def test_invalid_bank_amounts_are_rejected_without_rounding(value):
    with pytest.raises(ValueError): bank_amount(value)


@pytest.mark.parametrize('value,expected', [('0.01', '0.01'), ('123.40', '123.40'),
    ('99999999999.99', '99999999999.99'), (400, '400.00')])
def test_bank_amounts_are_exact_decimal_values(value, expected):
    assert format(bank_amount(value), '.2f') == expected


def test_contract_preserves_zeroes_and_bank_length_limit():
    assert contract_number(BillingStudent(id=1, student_name='Synthetic', home_location='MSA', school_student_id='002100')) == 'MSA0002100'
    assert len(contract_number(BillingStudent(id=1, student_name='Synthetic', home_location='MSA', school_student_id='1' * 16))) == 20


def test_large_decimal_workbook_cells_preserve_exact_bank_precision():
    from billing.bank_export import workbook_bytes
    from billing.bank_export_schemas import BankExportPreview, BankExportRow
    student = BillingStudent(id=1, student_name='Synthetic', home_location='MSA', school_student_id='1')
    review = BankExportPreview(collection_month='2026-11', branch='MSA', payment_count=2,
        total_amount='100000000000.00', content_hash='a' * 64, filename='test.xlsx',
        rows=[BankExportRow(payment_id='a', payment_version=1, student=student,
                            contract_no='MSA01', amount='99999999999.99'),
              BankExportRow(payment_id='b', payment_version=1, student=student,
                            contract_no='MSA02', amount='0.01')], issues=[], can_download=True)
    content = workbook_bytes(review)
    with ZipFile(BytesIO(content)) as archive:
        from xml.etree import ElementTree
        root = ElementTree.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        namespace = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        cells = {cell.attrib['r']: cell for cell in root.findall('m:sheetData/m:row/m:c', namespace)}
        exact = [cells[f'B{row}'].find('m:v', namespace).text for row in (2, 3)]
        assert exact == ['99999999999.99', '0.01']
        assert sum(Decimal(value) for value in exact) == Decimal(review.total_amount)
        assert all(cells[f'B{row}'].attrib['t'] == 'n' for row in (2, 3))
    sheet = load_workbook(BytesIO(content)).active
    assert str(sheet['B2'].value) == '99999999999.99'
    assert sheet['B2'].number_format == 'General'

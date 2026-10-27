"""Independent bank-file acceptance checks with invented SQLite evidence."""
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree

import pytest
from openpyxl import load_workbook
from sqlalchemy import event
from sqlalchemy.engine import Engine

from billing.models import BillingPayment, BillingPreference
from models import Student
from tests.test_billing_autopay_verification import (
    NOW, add_source, draft_env, payment, record_paid, records, run, workflow_env,
)


@pytest.fixture(autouse=True)
def reject_non_sqlite_connections():
    def reject(dialect, *_args):
        assert dialect.name == 'sqlite', 'Bank export verification refuses external database connections.'
    event.listen(Engine, 'do_connect', reject)
    yield
    event.remove(Engine, 'do_connect', reject)


@pytest.fixture
def bank_records(records):
    client, db, actor = records
    db.get(Student, 101).school_student_id = '2100'
    db.commit()
    return client, db, actor


def preview(client, branch='MSA', **changes):
    result = client.post('/api/billing/autopay/export-preview', json={
        'collection_month': '2026-11', 'branch': branch, **changes,
    })
    assert result.status_code == 200, result.text
    return result.json()


def export(client, reviewed, **changes):
    return client.post('/api/billing/autopay/export', json={
        'collection_month': reviewed['collection_month'], 'branch': reviewed['branch'],
        'expected_content_hash': reviewed['content_hash'], **changes,
    })


def assert_exact_workbook(response, reviewed):
    assert response.status_code == 200, response.text if response.status_code != 200 else ''
    assert response.content[:4] == b'PK\x03\x04'
    assert response.headers['x-csm-export-total'] == reviewed['total_amount']
    assert response.headers['x-csm-export-count'] == str(reviewed['payment_count'])
    assert response.headers['x-csm-export-hash'] == reviewed['content_hash']
    assert reviewed['filename'].isascii() and reviewed['filename'].endswith('.xlsx')
    assert reviewed['filename'] in response.headers['content-disposition']
    assert response.headers['content-type'].startswith('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert len(workbook.worksheets) == 1
    sheet = workbook.active
    assert sheet.max_row == reviewed['payment_count'] + 1 and sheet.max_column == 3
    assert [cell.value for cell in sheet[1]] == ['Contract No', 'Amount', 'Remark']
    total = Decimal('0')
    for index, expected in enumerate(reviewed['rows'], start=2):
        contract, amount, remark = sheet[index]
        assert contract.value == expected['contract_no']
        assert contract.data_type == 's' and contract.number_format == '@'
        assert amount.data_type == 'n' and amount.number_format == 'General'
        assert Decimal(str(amount.value)) == Decimal(expected['amount'])
        assert remark.value in (None, '')
        assert all(cell.data_type != 'f' and cell.hyperlink is None and cell.comment is None for cell in sheet[index])
        total += Decimal(str(amount.value))
    assert total == Decimal(reviewed['total_amount'])
    with ZipFile(BytesIO(response.content)) as archive:
        assert not any('externalLink' in name or 'vbaProject' in name for name in archive.namelist())
        xml = ElementTree.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        raw_values = [Decimal(cell.find('m:v', ns).text) for cell in xml.findall('.//m:c', ns)
                      if cell.attrib['r'].startswith('B') and cell.find('m:v', ns) is not None]
        assert sum(raw_values, Decimal('0')) == Decimal(response.headers['x-csm-export-total'])
    return sheet


def test_profile_branch_controls_the_file_independently_of_teaching_branch(bank_records):
    client, db, _ = bank_records
    db.get(Student, 101).home_location = 'MSB'
    db.commit()
    msa_taught = payment(client)
    add_source(db, student_id=102, source_id=21, location='MSB', school_id='0027')
    db.get(Student, 102).home_location = 'MSA'
    db.commit()
    msb_taught = payment(client, student_id=102, source_id=21, key='profile-cross-branch')
    teaching_list = run(client, location='MSA')
    assert {row['payment_id'] for row in teaching_list['rows'] if row['stage'] == 'ready'} == {msa_taught['id']}
    msa = preview(client, 'MSA')
    msb = preview(client, 'MSB')
    assert [row['payment_id'] for row in msa['rows']] == [msb_taught['id']]
    assert [row['payment_id'] for row in msb['rows']] == [msa_taught['id']]
    assert msa['rows'][0]['contract_no'] == 'MSA00027'
    assert msb['rows'][0]['contract_no'] == 'MSB02100'
    assert msa['total_amount'] == msb['total_amount'] == '1600.00'
    assert msa['filename'] != msb['filename'] and msa['content_hash'] != msb['content_hash']
    assert_exact_workbook(export(client, msa), msa)
    assert_exact_workbook(export(client, msb), msb)


def test_all_ready_rows_have_one_exact_total_and_no_sample_or_total_row(bank_records):
    client, db, _ = bank_records
    payment(client)
    add_source(db, student_id=102, source_id=21, location='MSA', school_id='27')
    payment(client, student_id=102, source_id=21, key='second-bank-row')
    add_source(db, student_id=103, source_id=31, location='MSA', school_id='28')
    payment(client, student_id=103, source_id=31, key='bank-unsent', sent=False)
    add_source(db, student_id=104, source_id=41, location='MSA', school_id='29')
    payment(client, student_id=104, source_id=41, key='bank-draft', draft_only=True)
    reviewed = preview(client)
    assert reviewed['can_download'] and not reviewed['issues']
    assert reviewed['payment_count'] == len(reviewed['rows']) == 2
    assert reviewed['total_amount'] == '3200.00'
    assert [row['student']['school_student_id'] for row in reviewed['rows']] == ['27', '2100']
    assert_exact_workbook(export(client, reviewed), reviewed)


def test_export_preview_and_download_do_not_mutate_any_financial_state(bank_records):
    client, db, _ = bank_records
    value = payment(client)
    before = deepcopy(db.get(BillingPayment, value['id']).__dict__)
    statements = []
    def only_select(_connection, _cursor, statement, *_args):
        statements.append(statement)
        assert statement.lstrip().upper().startswith('SELECT'), statement
    event.listen(db.bind, 'before_cursor_execute', only_select)
    try:
        reviewed = preview(client)
        response = export(client, reviewed)
        assert_exact_workbook(response, reviewed)
    finally:
        event.remove(db.bind, 'before_cursor_execute', only_select)
    db.expire_all()
    after = db.get(BillingPayment, value['id'])
    assert after.status == 'Pending Payment' and after.version == before['version']
    assert after.snapshot == before['snapshot'] and after.paid_at is None
    assert statements


@pytest.mark.parametrize('change', ['paid', 'contract', 'profile_branch', 'mandate', 'method', 'preference_version', 'payment_version', 'amount', 'new_ready'])
def test_changed_fresh_export_evidence_requires_a_new_review(bank_records, change):
    client, db, _ = bank_records
    value = payment(client)
    reviewed = preview(client)
    if change == 'paid':
        record_paid(client, value)
    elif change == 'contract':
        db.get(Student, 101).school_student_id = '2101'
    elif change == 'profile_branch':
        db.get(Student, 101).home_location = 'MSB'
    elif change == 'payment_version':
        db.get(BillingPayment, value['id']).version += 1
    elif change in ('mandate', 'method', 'preference_version'):
        preference = db.query(BillingPreference).one()
        if change == 'mandate':
            preference.mandate_ready = False
        elif change == 'method':
            preference.payment_method = 'manual'
        preference.version += 1
    elif change == 'amount':
        stored = db.get(BillingPayment, value['id'])
        snapshot = deepcopy(stored.snapshot)
        snapshot['review']['period']['total_fee'] += 1
        stored.snapshot = snapshot
    else:
        add_source(db, student_id=102, source_id=21, location='MSA', school_id='27')
        payment(client, student_id=102, source_id=21, key='new-bank-ready')
    db.commit()
    rejected = export(client, reviewed)
    assert rejected.status_code == 409, rejected.text
    assert not rejected.content.startswith(b'PK')
    fresh = preview(client)
    assert fresh['content_hash'] != reviewed['content_hash']


@pytest.mark.parametrize('school_id', [None, '', '=1+1', '中學2100', '21 00', 'A'*17])
def test_invalid_contract_cannot_silently_drop_one_student_or_create_a_formula(bank_records, school_id):
    client, db, _ = bank_records
    payment(client)
    add_source(db, student_id=102, source_id=21, location='MSA', school_id='27')
    payment(client, student_id=102, source_id=21, key='valid-bank-neighbour')
    db.get(Student, 101).school_student_id = school_id
    db.commit()
    reviewed = preview(client)
    assert not reviewed['can_download'] and reviewed['issues']
    assert export(client, reviewed).status_code == 422


@pytest.mark.parametrize('school_id,contract', [('A'*16, 'MSA0'+'A'*16), ('02100', 'MSA002100'), ('A-2100', 'MSA0A-2100')])
def test_valid_bank_contract_is_exact_text_with_internal_zeros_preserved(bank_records, school_id, contract):
    client, db, _ = bank_records
    payment(client)
    db.get(Student, 101).school_student_id = school_id
    db.commit()
    reviewed = preview(client)
    assert reviewed['can_download'] and reviewed['rows'][0]['contract_no'] == contract
    assert_exact_workbook(export(client, reviewed), reviewed)


def test_duplicate_contract_blocks_the_whole_branch_file(bank_records):
    client, db, _ = bank_records
    payment(client)
    add_source(db, student_id=102, source_id=21, location='MSA', school_id='2100')
    payment(client, student_id=102, source_id=21, key='duplicate-bank-contract')
    reviewed = preview(client)
    assert not reviewed['can_download'] and reviewed['issues']
    assert export(client, reviewed).status_code == 422


@pytest.mark.parametrize('branch', [None, 'Various'])
def test_unknown_profile_branch_blocks_both_files_for_review(bank_records, branch):
    client, db, _ = bank_records
    payment(client)
    db.get(Student, 101).home_location = branch
    db.commit()
    for known in ['MSA', 'MSB']:
        reviewed = preview(client, known)
        assert not reviewed['can_download'] and reviewed['issues']
        assert export(client, reviewed).status_code == 422


def test_invalid_contract_in_the_other_known_branch_does_not_block_this_file(bank_records):
    client, db, _ = bank_records
    payment(client)
    add_source(db, student_id=102, source_id=21, location='MSB', school_id='=BAD')
    db.get(Student, 102).home_location = 'MSB'
    db.commit()
    payment(client, student_id=102, source_id=21, key='other-branch-invalid')
    reviewed = preview(client, 'MSA')
    assert reviewed['can_download'] and reviewed['payment_count'] == 1
    assert_exact_workbook(export(client, reviewed), reviewed)
    assert not preview(client, 'MSB')['can_download']


@pytest.mark.parametrize('role,effective,allowed', [
    ('Admin', None, True), ('Super Admin', None, True), ('Super Admin', 'Admin', True),
    ('Supervisor', None, False), ('Tutor', None, False), ('Guest', None, False),
    ('Super Admin', 'Tutor', False), ('Super Admin', 'Supervisor', False),
    ('Tutor', 'Admin', False), ('Supervisor', 'Super Admin', False),
])
def test_both_export_routes_require_effective_admin_without_header_escape(bank_records, role, effective, allowed):
    client, _db, actor = bank_records
    payment(client)
    reviewed = preview(client)
    actor.role = role
    headers = {'X-Effective-Role': effective} if effective else {}
    result = client.post('/api/billing/autopay/export-preview', headers=headers,
        json={'collection_month': '2026-11', 'branch': 'MSA'})
    downloaded = client.post('/api/billing/autopay/export', headers=headers,
        json={'collection_month': '2026-11', 'branch': 'MSA', 'expected_content_hash': reviewed['content_hash']})
    assert result.status_code == downloaded.status_code == (200 if allowed else 403)
    if not allowed:
        assert not downloaded.content.startswith(b'PK')


def test_storage_unavailable_never_queries_missing_188_or_189_tables(bank_records, monkeypatch):
    client, db, _ = bank_records
    monkeypatch.setenv('BILLING_WORKFLOW_ENABLED', 'false')
    monkeypatch.setenv('BILLING_WORKFLOW_STORAGE_ENABLED', 'false')
    queries = []
    def no_new_storage(_connection, _cursor, statement, *_args):
        text = statement.lower()
        queries.append(text)
        assert not any(name in text for name in ['billing_accounts', 'billing_payments', 'billing_payment_lines',
            'billing_lesson_claims', 'billing_preferences', 'student_participation', 'makeup_entitlement'])
    event.listen(db.bind, 'before_cursor_execute', no_new_storage)
    try:
        reviewed = preview(client)
        assert not reviewed['can_download'] and reviewed['payment_count'] == 0
        response = export(client, reviewed)
        assert response.status_code == 503, response.text
    finally:
        event.remove(db.bind, 'before_cursor_execute', no_new_storage)


@pytest.mark.parametrize('body', [
    {'collection_month': '2026-10', 'branch': 'MSA'},
    {'collection_month': '2027-07', 'branch': 'MSA'},
    {'collection_month': '2026-11', 'branch': 'Various'},
    {'collection_month': '2026-11', 'branch': 'MSA', 'expected_content_hash': ''},
])
def test_invalid_month_branch_or_hash_cannot_return_a_workbook(bank_records, body):
    client, _db, _ = bank_records
    response = client.post('/api/billing/autopay/export', json=body)
    assert response.status_code == 422, response.text
    assert not response.content.startswith(b'PK')


@pytest.mark.parametrize('value', [0, -1, Decimal('100000000000'), Decimal('1.001'),
    Decimal('NaN'), Decimal('Infinity'), True, None, '1,200.00', '$1200'])
def test_bank_amount_rejects_unsupported_values_without_rounding(value):
    from billing.bank_export import bank_amount
    with pytest.raises(ValueError):
        bank_amount(value)


@pytest.mark.parametrize('value,expected', [(Decimal('0.01'), Decimal('0.01')),
    (Decimal('1.2300'), Decimal('1.23')), (Decimal('99999999999.99'), Decimal('99999999999.99'))])
def test_exact_decimal_bank_values_survive_validation(value, expected):
    from billing.bank_export import bank_amount
    assert bank_amount(value) == expected


def test_fractional_workbook_values_sum_exactly_in_the_actual_xml():
    from billing.bank_export import workbook_bytes
    from billing.bank_export_schemas import BankExportPreview, BankExportRow
    from billing.schemas import BillingStudent
    amounts = ['0.10', '0.20', '99999999999.99']
    reviewed = BankExportPreview(collection_month='2026-11', branch='MSA',
        payment_count=3, total_amount='100000000000.29', content_hash='a'*64,
        filename='autopay-MSA-2026-11-test.xlsx', issues=[], can_download=True,
        rows=[BankExportRow(payment_id=f'fraction-{index}', payment_version=1,
            student=BillingStudent(id=index, student_name=f'Synthetic {index}'),
            contract_no=f'MSA000{index}', amount=amount) for index, amount in enumerate(amounts, 1)])
    with ZipFile(BytesIO(workbook_bytes(reviewed))) as archive:
        xml = ElementTree.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        values = [Decimal(cell.find('m:v', ns).text) for cell in xml.findall('.//m:c', ns)
            if cell.attrib['r'].startswith('B') and cell.find('m:v', ns) is not None]
    assert values == [Decimal(amount) for amount in amounts]
    assert sum(values, Decimal('0')) == Decimal(reviewed.total_amount)


@pytest.mark.parametrize('field,value', [('search', 'nobody'), ('offset', 999), ('limit', 1),
    ('stage', 'skipped'), ('location', 'MSB'), ('tutor_id', 999)])
def test_list_filters_cannot_be_smuggled_into_the_bank_export_scope(bank_records, field, value):
    client, _db, _ = bank_records
    payment(client)
    response = client.post('/api/billing/autopay/export-preview', json={
        'collection_month': '2026-11', 'branch': 'MSA', field: value})
    assert response.status_code == 422


def test_empty_review_never_downloads_an_empty_bank_file(bank_records):
    client, _db, _ = bank_records
    reviewed = preview(client)
    assert not reviewed['can_download'] and reviewed['total_amount'] == '0.00'
    assert reviewed['payment_count'] == 0 and reviewed['rows'] == []
    assert export(client, reviewed).status_code == 422

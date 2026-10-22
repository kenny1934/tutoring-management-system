"""Coupon source and commitment invariants using invented isolated records."""
from copy import deepcopy
from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook
from sqlalchemy import event
from sqlalchemy.engine import Engine

from billing import coupon_accounting as accounting
from billing.coupon_models import CouponAccount, CouponImport, CouponObservation, CouponUse
from billing.models import BillingPayment, BillingPaymentLine
from models import Enrollment, Student
from tests.test_billing_drafts import env as draft_env
from tests.test_billing_workflow import env as workflow_env, confirmed, coupon_payload, pay_payload, second_class


@pytest.fixture(autouse=True)
def sqlite_only():
    def reject(dialect, *_args):
        assert dialect.name == 'sqlite', 'Coupon checks refuse external database connections.'
    event.listen(Engine, 'do_connect', reject)
    yield
    event.remove(Engine, 'do_connect', reject)


@pytest.fixture
def env(workflow_env, monkeypatch):
    from routers import billing_coupons
    monkeypatch.setenv('BILLING_COUPON_ENABLED', 'true')
    monkeypatch.setenv('BILLING_COUPON_STORAGE_ENABLED', 'true')
    monkeypatch.setattr(billing_coupons, 'today', lambda: date(2026, 10, 17))
    client, db, actor = workflow_env
    db.get(Student, 101).school_student_id = '0021'
    db.commit()
    client.app.include_router(billing_coupons.router, prefix='/api')
    return client, db, actor


def workbook(rows):
    book = Workbook()
    sheet = book.active
    sheet.append(['ID#', *['Unused'] * 9, 'Coupon'])
    for company, count in rows:
        sheet.append([company, *[''] * 9, count])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def balance(client):
    response = client.get('/api/billing/students/101/coupons')
    assert response.status_code == 200, response.text
    return response.json()


def preview(client, rows, observed_on='2026-10-15'):
    response = client.post('/api/billing/coupon-imports/preview', data={'branch': 'MSA', 'observed_on': observed_on},
        files={'file': ('invented-kms.xls', workbook(rows), 'application/octet-stream')})
    assert response.status_code == 200, response.text
    return response.json()


def apply(client, reviewed, *, key='import-source', **extra):
    return client.post(f"/api/billing/coupon-imports/{reviewed['id']}/confirm", json={
        'expected_review_hash': reviewed['review_hash'], 'idempotency_key': key, **extra})


def observed(client, count=1, observed_on='2026-10-15', key='import-source'):
    source = preview(client, [('MSA0021', count)], observed_on)
    result = apply(client, source, key=key)
    assert result.status_code == 200, result.text
    return result.json()


def checked(client, groups=None, reconciliations=None, key='check-source'):
    current = balance(client)
    payload = {'expected_version': current['version'], 'observation_id': current['observation']['id'],
        'checked_at': '2026-10-17T01:00:00Z',
        'groups': groups if groups is not None else [{'kind': 'ordinary', 'value': 300, 'expires_on': '2027-06-30', 'quantity': 1}],
        'reconciliations': reconciliations or [], 'reason': 'Checked the actual KMS coupon details.'}
    reviewed = client.post('/api/billing/students/101/coupons/review', json=payload)
    assert reviewed.status_code == 200, reviewed.text
    response = client.post('/api/billing/students/101/coupons/verify', json={**payload,
        'expected_review_hash': reviewed.json()['review_hash'], 'idempotency_key': key})
    return response, reviewed.json(), payload


def payment(client):
    return confirmed(client, payload=coupon_payload(client))[2]


def paid(client, value):
    response = client.post(f"/api/billing/payments/{value['id']}/paid", json=pay_payload(value))
    assert response.status_code == 200, response.text
    return response.json()


def receipt(client, value, **extra):
    return client.post(f"/api/billing/payments/{value['id']}/receipt", json={
        'expected_version': value['version'], 'idempotency_key': 'receipt-source',
        'receipt_reference': 'SYNTHETIC-RECEIPT', **extra})


def test_observation_is_separate_and_only_checked_details_make_stock_usable(env):
    client, db, _ = env
    observed(client)
    current = balance(client)
    assert current['observation']['remaining_count'] == 1 and current['usable_count'] is None
    assert current['observation']['observed_on'] == '2026-10-15'
    assert current['observation']['source_sha256']
    response, _, _ = checked(client)
    assert response.status_code == 200, response.text
    assert response.json()['usable_count'] == 1 and response.json()['checked_by'] == 'Draft Admin'
    assert db.query(CouponObservation).one().remaining_count == 1


def test_fresh_zero_without_unresolved_uses_needs_no_value_or_expiry_check(env):
    client, _db, _ = env
    observed(client, 0)
    current = balance(client)
    assert current['usable_count'] == 0 and current['groups'] == [] and current['checked_at'] is None
    assert not current['issues']


def test_zero_observation_does_not_clear_an_unresolved_reservation(env):
    client, _db, _ = env
    observed(client)
    assert checked(client)[0].status_code == 200
    payment(client)
    observed(client, 0, '2026-10-17', 'zero-with-unpaid-commitment')
    current = balance(client)
    assert current['observation']['remaining_count'] == 0 and current['usable_count'] is None
    assert current['reserved_count'] == 1 and not current['uses'][0]['accounted_in_observation']


def test_one_combined_coupon_stays_consumed_after_receipt_against_old_observation(env):
    client, db, _ = env
    observed(client)
    assert checked(client)[0].status_code == 200
    second_class(db)
    value = confirmed(client, payload=coupon_payload(client, ids=(11, 12)))[2]
    assert len(value['lines']) == 2
    assert balance(client)['usable_count'] == 0 and db.query(CouponUse).count() == 1
    value = paid(client, value)
    assert receipt(client, value).status_code == 422
    completed = receipt(client, value, coupon_selected=True, coupon_used_on='2026-10-17')
    assert completed.status_code == 200, completed.text
    current = balance(client)
    assert current['observation']['remaining_count'] == 1 and current['usable_count'] == 0
    assert current['used_awaiting_sync_count'] == 1 and current['receipt_pending_count'] == 0
    assert db.query(CouponUse).one().state == 'receipt_recorded'
    assert sum(line.discount_amount for line in db.query(BillingPaymentLine)) == 300


def test_fresh_count_drop_never_reconciles_a_specific_use_without_review(env):
    client, db, _ = env
    observed(client)
    assert checked(client)[0].status_code == 200
    value = paid(client, payment(client))
    assert receipt(client, value, coupon_selected=True).status_code == 200
    observed(client, 0, '2026-10-17', 'new-source')
    current = balance(client)
    assert current['usable_count'] is None and current['used_awaiting_sync_count'] == 1
    use = current['uses'][0]
    assert not use['accounted_in_observation']
    response, review, _ = checked(client, groups=[], reconciliations=[{
        'use_id': use['id'], 'included_in_observation': True}], key='checked-new-source')
    assert response.status_code == 200, review
    assert response.json()['usable_count'] == 0 and response.json()['used_awaiting_sync_count'] == 0
    assert db.query(CouponObservation).count() == 2


def test_receipt_after_import_does_not_double_subtract_a_reconciled_use(env):
    client, _db, _ = env
    observed(client, 2)
    groups = [{'kind': 'ordinary', 'value': 300, 'expires_on': '2027-06-30', 'quantity': 2}]
    assert checked(client, groups)[0].status_code == 200
    value = paid(client, payment(client))
    observed(client, 1, '2026-10-17', 'after-kms-redeemed')
    use_id = balance(client)['uses'][0]['id']
    response, _, _ = checked(client, reconciliations=[{'use_id': use_id, 'included_in_observation': True}], key='includes-use')
    assert response.status_code == 200, response.text
    assert response.json()['usable_count'] == 1 and response.json()['receipt_pending_count'] == 1
    assert receipt(client, value, coupon_selected=True, coupon_used_on='2026-10-16').status_code == 200
    assert balance(client)['usable_count'] == 1


def test_mixed_groups_cannot_fund_a_different_value(env):
    client, db, _ = env
    observed(client, 2)
    groups = [{'kind': 'ordinary', 'value': 200, 'expires_on': '2027-06-30', 'quantity': 1},
              {'kind': 'ordinary', 'value': 300, 'expires_on': '2027-06-30', 'quantity': 1}]
    assert checked(client, groups)[0].status_code == 200
    payment(client)
    current = balance(client)
    assert {g['value']: g['usable_count'] for g in current['groups']} == {200: 1, 300: 0}
    assert current['usable_count'] == 1 and db.query(CouponUse).count() == 1


@pytest.mark.parametrize('count', [-1, 1.2, 'abc', '=1+1', True])
def test_invalid_source_count_never_becomes_zero_or_changes_source(env, count):
    client, db, _ = env
    source = preview(client, [('MSA0021', count)])
    assert source['invalid_count'] == 1 and not source['can_confirm']
    assert source['rows'][0]['remaining_count'] is None
    assert apply(client, source).status_code == 409
    assert db.query(CouponObservation).count() == 0


def test_exact_school_id_and_unmatched_require_explicit_exclusion(env):
    client, db, _ = env
    source = preview(client, [('MSA0021', 1), ('MSA21', 7)])
    assert source['changed_count'] == 1 and source['unmatched_count'] == 1
    assert source['rows'][0]['student']['id'] == 101 and source['rows'][1]['student'] is None
    assert apply(client, source).status_code == 409
    assert apply(client, source, acknowledge_unmatched=True).status_code == 200
    assert db.query(CouponObservation).one().remaining_count == 1


def test_same_day_conflicting_source_requires_latest_export_attestation(env):
    client, db, _ = env
    observed(client, 1)
    source = preview(client, [('MSA0021', 2)])
    assert any(i['code'] == 'SAME_DAY_SOURCE_REVIEW' for i in source['rows'][0]['issues'])
    assert source['can_confirm'] and apply(client, source, key='same-day').status_code == 409
    assert apply(client, source, key='same-day-ack', acknowledge_same_day=True).status_code == 200
    assert db.get(CouponObservation, db.get(CouponAccount, 101).observation_id).remaining_count == 2


def test_same_file_cannot_become_newer_by_changing_observation_date(env):
    client, db, _ = env
    content = workbook([('MSA0021', 1)])
    def upload(day):
        response = client.post('/api/billing/coupon-imports/preview', data={'branch': 'MSA', 'observed_on': day},
            files={'file': ('same-source.xls', content, 'application/octet-stream')})
        assert response.status_code == 200, response.text
        return response.json()
    first = upload('2026-10-15')
    assert apply(client, first).status_code == 200
    later = upload('2026-10-17')
    assert not later['can_confirm'] and any(i['code'] == 'SOURCE_DATE_CHANGED' for i in later['issues'])
    assert apply(client, later, key='changed-as-of').status_code == 409
    assert db.query(CouponObservation).count() == 1


def test_off_helper_has_no_191_queries_or_changed_legacy_snapshot_shape(env, monkeypatch):
    client, db, _ = env
    monkeypatch.setenv('BILLING_COUPON_ENABLED', 'false')
    monkeypatch.setenv('BILLING_COUPON_STORAGE_ENABLED', 'false')
    def no_191(_conn, _cursor, sql, *_args):
        assert 'billing_coupon_' not in sql.lower()
    event.listen(db.bind, 'before_cursor_execute', no_191)
    try:
        assert accounting.balances_for_students(db, [101]) == {}
        assert not balance(client)['storage_available']
    finally:
        event.remove(db.bind, 'before_cursor_execute', no_191)
    from billing.schemas import ImportedCouponSnapshot, BillingPaidRequest
    assert 'accounting' not in ImportedCouponSnapshot(remaining_count=0, cached_value=None, synced_at_utc=None).model_dump()
    data = BillingPaidRequest(expected_version=1, idempotency_key='old-hash', payment_date=date(2026, 10, 17), amount=400, method='cash')
    assert not {'coupon_selected', 'coupon_used_on', 'acknowledge_duplicate_receipt'} & data.model_dump().keys()


def test_cancelled_consumed_legacy_use_does_not_replenish_old_source(env):
    client, db, actor = env
    observed(client)
    assert checked(client)[0].status_code == 200
    use = CouponUse(id='synthetic-legacy-use', student_id=101, legacy_enrollment_id=11, value=300,
        group_key=balance(client)['groups'][0]['key'], verification_id=db.get(CouponAccount, 101).verification_id,
        expires_on=date(2027, 6, 30), state='receipt_recorded', receipt_reference='SYNTHETIC-OLD')
    db.add(use)
    db.get(Enrollment, 11).payment_status = 'Cancelled'
    db.commit()
    current = balance(client)
    assert current['usable_count'] is None and current['used_awaiting_sync_count'] == 1
    assert current['uses'][0]['state'] == 'receipt_recorded'
    assert any(i['code'] == 'COUPON_SOURCE_REVIEW' for i in current['issues'])


def test_known_kms_metadata_and_verified_footer_are_not_student_records():
    from billing.coupon_imports import KMS_HEADERS, parse_workbook
    book = Workbook()
    sheet = book.active
    sheet.append(['MSA / 2027-01-01'])  # This cutoff is not the observed-on date.
    sheet.append(KMS_HEADERS)
    sheet.append(['MSA0021', 'Synthetic', '', '', '', '', '', '', '', '', 1])
    sheet.append(['MSA0022', 'Synthetic', '', '', '', '', '', '', '', '', 2])
    sheet.append(['', '', '', '', '', '', '', '12', '3', '--', '3', '--', '', ''])
    output = BytesIO()
    book.save(output)
    parsed = parse_workbook(output.getvalue())
    assert [(r['row_number'], r['company_id'], r['remaining_count']) for r in parsed] == [(3, 'MSA0021', 1), (4, 'MSA0022', 2)]
    sheet.cell(5, 1, 'MSA0023')
    output = BytesIO()
    book.save(output)
    parsed = parse_workbook(output.getvalue())
    assert len(parsed) == 3 and parsed[-1]['company_id'] == 'MSA0023'


@pytest.mark.parametrize('corruption', ['foreign_review', 'invalid_review', 'foreign_line'])
def test_payment_detail_holds_foreign_or_malformed_evidence_without_echoing_it(env, corruption):
    client, db, _ = env
    value = confirmed(client)[2]
    stored = db.get(BillingPayment, value['id'])
    if corruption == 'foreign_line':
        enrollment = db.get(Enrollment, value['lines'][0]['enrollment_id'])
        enrollment.student_id = 102
    else:
        snapshot = deepcopy(stored.snapshot)
        snapshot['review']['student']['student_name'] = 'PRIVATE_OTHER_STUDENT'
        if corruption == 'foreign_review':
            snapshot['review']['student']['id'] = 102
        else:
            snapshot['review']['student']['id'] = {'malformed': 'PRIVATE_OTHER_STUDENT'}
        stored.snapshot = snapshot
    db.commit()
    response = client.get(f"/api/billing/payments/{value['id']}")
    assert response.status_code == 409 and 'PRIVATE_OTHER_STUDENT' not in response.text


@pytest.mark.parametrize('problem', ['forged_value', 'missing_groups', 'known_zero', 'forged_expiry'])
def test_draft_quote_save_and_copy_require_authoritative_available_group(env, problem):
    client, db, _ = env
    observed(client, 0 if problem == 'known_zero' else 1)
    if problem not in {'known_zero', 'missing_groups'}:
        assert checked(client)[0].status_code == 200
    request = coupon_payload(client, value=200 if problem == 'forged_value' else 300)
    if problem == 'forged_expiry':
        request['scenario']['coupon']['expires_on'] = '2027-05-31'
    request['confirmed_checks'].append('COUPON_STOCK_REVIEW')
    assert client.post('/api/billing/students/101/draft-review', json=request).status_code == 422
    request['confirmed_checks'].remove('COUPON_STOCK_REVIEW')
    reviewed = client.post('/api/billing/students/101/draft-review', json=request)
    assert reviewed.status_code == 200, reviewed.text
    quote = reviewed.json()
    assert 'COUPON_STOCK_REVIEW' in {i['code'] for i in quote['issues']}
    assert quote['message']['zh'] is None and quote['message']['en'] is None
    saved = client.post('/api/billing/students/101/drafts', json={
        'review': request, 'expected_review_hash': quote['review_hash'], 'idempotency_key': 'blocked-coupon-draft'})
    assert saved.status_code == 200, saved.text
    assert not saved.json()['can_copy']
    fresh = client.get(f"/api/billing/drafts/{saved.json()['id']}")
    assert fresh.status_code == 200 and not fresh.json()['can_copy']
    assert fresh.json()['review']['message']['zh'] is None and fresh.json()['review']['message']['en'] is None
    assert db.query(BillingPayment).count() == db.query(CouponUse).count() == 0


def test_known_available_group_still_allows_review_save_and_copy(env):
    client, _db, _ = env
    observed(client)
    assert checked(client)[0].status_code == 200
    request = coupon_payload(client)
    reviewed = client.post('/api/billing/students/101/draft-review', json=request)
    assert reviewed.status_code == 200, reviewed.text
    quote = reviewed.json()
    assert 'COUPON_STOCK_REVIEW' not in {i['code'] for i in quote['issues']}
    assert quote['message']['zh'] and quote['message']['en']
    saved = client.post('/api/billing/students/101/drafts', json={
        'review': request, 'expected_review_hash': quote['review_hash'], 'idempotency_key': 'valid-coupon-draft'})
    assert saved.status_code == 200 and saved.json()['can_copy'], saved.text
    fresh = client.get(f"/api/billing/drafts/{saved.json()['id']}")
    assert fresh.status_code == 200 and fresh.json()['can_copy']
